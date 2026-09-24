from pathlib import Path
from datetime import datetime
import re
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parents[1]
SRC = BASE / "docs" / "architecture" / "system.md"
EXPORT_DIR = BASE / "docs" / "exports"
NAVY, PALE, BORDER = "17365D", "F4F8FB", "D9D9D9"
BLACK, WHITE, GRAY = "000000", "FFFFFF", "666666"

def font(run, name="Arial", size=10.5, bold=False, color=BLACK):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    for key in ("ascii", "hAnsi", "eastAsia"):
        rpr.rFonts.set(qn("w:" + key), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    node = tcpr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tcpr.append(node)
    node.set(qn("w:fill"), fill)

def borders(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    group = tcpr.first_child_found_in("w:tcBorders")
    if group is None:
        group = OxmlElement("w:tcBorders")
        tcpr.append(group)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = group.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            group.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), BORDER)

def margins(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    group = tcpr.first_child_found_in("w:tcMar")
    if group is None:
        group = OxmlElement("w:tcMar")
        tcpr.append(group)
    for edge in ("top", "start", "bottom", "end"):
        node = group.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            group.append(node)
        node.set(qn("w:w"), "110")
        node.set(qn("w:type"), "dxa")

def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    trpr.append(node)

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Страница ")
    font(run, size=8.5, color=GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])

def configure(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.7)
    section.left_margin, section.right_margin = Cm(2.0), Cm(1.7)
    add_page_number(section.footer.paragraphs[0])

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(13)

    for style_name, size, before, after in (
        ("Heading 1", 15, 15, 7),
        ("Heading 2", 12, 10, 5),
        ("Heading 3", 10.5, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

def add_text(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(4.5)
    p.paragraph_format.line_spacing = 1.12
    run = p.add_run(text)
    font(run)
    return p

def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.right_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.0
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    ppr.append(shd)
    run = p.add_run(text)
    font(run, name="Consolas", size=8.5)

def add_table(doc, rows):
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=1, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header = table.rows[0]
    repeat_header(header)
    for i in range(cols):
        value = rows[0][i] if i < len(rows[0]) else ""
        cell = header.cells[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, NAVY)
        margins(cell)
        borders(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        font(p.add_run(value), size=8.7, bold=True, color=WHITE)
    for ridx, values in enumerate(rows[1:]):
        row = table.add_row()
        for i in range(cols):
            value = values[i] if i < len(values) else ""
            cell = row.cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cell)
            borders(cell)
            if ridx % 2:
                shade(cell, PALE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.04
            font(p.add_run(value), size=8.6)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def parse_table(lines, start):
    rows, i = [], start
    while i < len(lines) and lines[i].strip().startswith("|"):
        parts = [p.strip() for p in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", p or "") for p in parts):
            rows.append(parts)
        i += 1
    return rows, i

def build():
    lines = SRC.read_text(encoding="utf-8-sig").splitlines()
    version = "unknown"
    if lines and lines[0] == "---":
        end = lines.index("---", 1)
        metadata = dict(line.split(":", 1) for line in lines[1:end] if ":" in line)
        version = metadata.get("version", "unknown").strip()
        if not re.fullmatch(r"[0-9]+\.[0-9]+", version):
            raise ValueError("Invalid document version")
        lines = lines[end + 1:]
    # Export the source body, not repository navigation or YAML metadata.
    first_heading = next(i for i, line in enumerate(lines) if line.startswith("# "))
    lines = lines[first_heading:]
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    output = EXPORT_DIR / f"architecture-spec-job-search-ai-agent-v{version}-{datetime.now():%Y%m%d-%H%M%S-%f}.docx"
    doc = Document()
    configure(doc)
    i, in_code, code = 0, False, []
    fence = chr(96) * 3

    while i < len(lines):
        raw, line = lines[i], lines[i].strip()
        if line.startswith(fence):
            if in_code:
                add_code(doc, "\n".join(code))
                code, in_code = [], False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code.append(raw)
            i += 1
            continue
        if not line:
            i += 1
            continue
        if line == "<!-- PAGEBREAK -->":
            doc.add_page_break()
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            rows, i = parse_table(lines, i)
            if rows:
                add_table(doc, rows)
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.add_run(line[2:])
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            add_text(doc, line[2:], style="List Bullet")
        elif re.match(r"^\d+\. ", line):
            add_text(doc, re.sub(r"^\d+\. ", "", line), style="List Number")
        else:
            p = add_text(doc, line)
            if "Веб сервис с backend" in line:
                for r in p.runs:
                    r.font.size = Pt(12)
                    r.font.color.rgb = RGBColor.from_string(GRAY)
                p.paragraph_format.space_after = Pt(16)
        i += 1

    props = doc.core_properties
    props.title = "Архитектурное техническое задание автономного ИИ агента поиска работы"
    props.subject = "Архитектура веб сервиса с backend на Python"
    props.author = ""
    props.keywords = "AI agent, job search, Python, FastAPI, Temporal, architecture"
    props.comments = ""
    doc.save(output)
    print(output)

if __name__ == "__main__":
    build()
