from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from cv_backend.services.document_extractor import UnsupportedDocumentError, UnsupportedDocumentTypeError, extract_document_text


def test_extracts_docx_paragraphs_without_retaining_upload():
    stream = BytesIO()
    document = Document()
    document.add_paragraph("Опыт работы")
    document.add_paragraph("Инженер")
    document.save(stream)
    assert extract_document_text("resume.docx", stream.getvalue()) == "Опыт работы\nИнженер"


def test_rejects_extension_signature_mismatch():
    with pytest.raises(UnsupportedDocumentError):
        extract_document_text("resume.pdf", b"PK not a PDF")


def test_rejects_unsupported_legacy_word_format_separately():
    with pytest.raises(UnsupportedDocumentTypeError):
        extract_document_text("resume.doc", b"legacy word bytes")


def test_rejects_empty_text_layer():
    with pytest.raises(UnsupportedDocumentError, match="text layer"):
        extract_document_text("scanned.pdf", minimal_pdf_without_text())


def test_extracts_pdf_text_layer():
    assert "Experience" in extract_document_text("resume.pdf", minimal_pdf_with_text())


def minimal_pdf_without_text() -> bytes:
    stream = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.write(stream)
    return stream.getvalue()


def minimal_pdf_with_text() -> bytes:
    stream = BytesIO()
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({
            NameObject("/F1"): DictionaryObject({
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            })
        })
    })
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 72 720 Td (Experience) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    writer.write(stream)
    return stream.getvalue()
