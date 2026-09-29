from uuid import UUID, uuid4
from io import BytesIO

from docx import Document
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from cv_backend.domain.candidate import SkillItemInput
from cv_backend.storage.models.candidate import CandidateItemModel, CandidateItemVersionModel
from cv_backend.storage.models.draft import DraftBlockModel, ResumeDraftModel
from cv_backend.storage.models.profile import ProfileItemSelectionModel, SpecializationProfileModel
from cv_backend.storage.repositories.candidate import CandidateRepository
TEST_USER_ID = UUID("00000000-0000-4000-8000-000000000001")


def test_user_can_create_and_list_multiple_profiles(client):
    first = client.post("/api/v1/profiles", json={"name": "AI Engineer"})
    second = client.post("/api/v1/profiles", json={"name": "QA Engineer"})
    assert first.status_code == 201
    assert second.status_code == 201
    listing = client.get("/api/v1/profiles")
    assert [profile["name"] for profile in listing.json()] == ["AI Engineer", "QA Engineer"]


def test_user_creates_empty_draft_and_adds_experience_as_its_own_block(client):
    response = client.post("/api/v1/resume-drafts")
    assert response.status_code == 201
    assert response.json()["state"] == "needs_user_review"
    assert response.json()["blocks"] == []
    added = client.post(
        f"/api/v1/resume-drafts/{response.json()['draft_id']}/experience-blocks",
        json={"text": "Python"},
    )
    assert added.status_code == 200
    assert added.json()["blocks"][0]["kind"] == "experience"
    assert added.json()["blocks"][0]["ordinal"] == 0
    assert client.post("/api/v1/resume-drafts/text", json={"text": "Python"}).status_code == 405
    assert client.get("/api/v1/candidate-base/items").json() == []


def test_profile_patch_does_not_change_another_profile(client):
    one = client.post("/api/v1/profiles", json={"name": "AI"}).json()
    two = client.post("/api/v1/profiles", json={"name": "QA"}).json()
    changed = client.patch(f"/api/v1/profiles/{one['id']}", json={"headline": "AI leader"})
    assert changed.status_code == 200
    untouched = client.get(f"/api/v1/profiles/{two['id']}")
    assert untouched.json()["headline"] is None


def test_unknown_profile_is_not_found(client):
    response = client.get(f"/api/v1/profiles/{uuid4()}")
    assert response.status_code == 404


def test_docx_import_creates_blocks_and_discards_original_upload(client):
    stream = BytesIO()
    document = Document()
    document.add_paragraph("Навыки")
    document.add_paragraph("Автоматизация")
    document.save(stream)
    response = client.post(
        "/api/v1/resume-drafts/file",
        files={"file": ("resume.docx", stream.getvalue(), "application/octet-stream")},
    )
    assert response.status_code == 201
    assert response.json()["blocks"][0]["text"] == "Автоматизация"
    assert "source_file" not in response.json()


def test_reimporting_identical_document_creates_another_draft(client):
    stream = BytesIO()
    document = Document()
    document.add_paragraph("Навыки")
    document.add_paragraph("Python")
    document.save(stream)
    payload = stream.getvalue()
    first = client.post(
        "/api/v1/resume-drafts/file",
        files={"file": ("resume.docx", payload, "application/octet-stream")},
    )
    second = client.post(
        "/api/v1/resume-drafts/file",
        files={"file": ("resume.docx", payload, "application/octet-stream")},
    )
    assert first.status_code == second.status_code == 201
    assert first.json()["draft_id"] != second.json()["draft_id"]


def test_upload_rejects_mismatched_extension_and_bytes(client):
    response = client.post(
        "/api/v1/resume-drafts/file",
        files={"file": ("resume.pdf", b"PK not PDF", "application/pdf")},
    )
    assert response.status_code == 422
    assert client.get("/api/v1/resume-drafts").json() == []


def test_upload_rejects_unsupported_extension_with_415(client, monkeypatch):
    response = client.post(
        "/api/v1/resume-drafts/file",
        files={"file": ("resume.doc", b"legacy format", "application/msword")},
    )
    assert response.status_code == 415


def test_upload_enforces_size_limit(client, monkeypatch):
    from cv_backend.api.routes import resume_drafts

    monkeypatch.setattr(resume_drafts, "MAX_UPLOAD_BYTES", 4)
    response = client.post(
        "/api/v1/resume-drafts/file",
        files={"file": ("resume.pdf", b"%PDF-1.4", "application/pdf")},
    )
    assert response.status_code == 413


def test_apply_idempotency_key_cannot_be_reused_for_another_block(client):
    imported = client.post("/api/v1/resume-drafts").json()
    imported = client.post(
        f"/api/v1/resume-drafts/{imported['draft_id']}/experience-blocks",
        json={"text": "Python and Git"},
    ).json()
    first_block = imported["blocks"][0]
    imported = client.post(
        f"/api/v1/resume-drafts/{imported['draft_id']}/experience-blocks",
        json={"text": "Git"},
    ).json()
    second_block = imported["blocks"][1]
    payload = {"items": [{"kind": "skill", "name": "Python"}], "idempotency_key": "same-key"}
    first = client.post(
        f"/api/v1/resume-drafts/{imported['draft_id']}/blocks/{first_block['id']}/apply", json=payload
    )
    conflict = client.post(
        f"/api/v1/resume-drafts/{imported['draft_id']}/blocks/{second_block['id']}/apply", json=payload
    )
    assert first.status_code == 200
    assert conflict.status_code == 409


def test_profile_delete_hides_profile_but_retains_profile_selection(client):
    item_id = None
    with Session(client.app.state.database_engine) as session:
        item_id = CandidateRepository(session).add_items(
            TEST_USER_ID, [SkillItemInput(name="Python", level="used")]
        )[0].id
        session.commit()

    profile = client.post("/api/v1/profiles", json={"name": "QA"}).json()
    selected = client.put(
        f"/api/v1/profiles/{profile['id']}/selections", json={"skill_ids": [str(item_id)]}
    )
    assert selected.status_code == 200

    assert client.delete(f"/api/v1/profiles/{profile['id']}").status_code == 204
    assert client.get(f"/api/v1/profiles/{profile['id']}").status_code == 404
    assert client.get("/api/v1/profiles").json() == []
    with Session(client.app.state.database_engine) as session:
        stored = session.get(SpecializationProfileModel, UUID(profile["id"]))
        assert stored is not None and stored.is_deleted and stored.deleted_at is not None
        assert session.scalar(
            select(func.count()).select_from(ProfileItemSelectionModel).where(
                ProfileItemSelectionModel.profile_id == UUID(profile["id"])
            )
        ) == 1


def test_candidate_item_delete_hides_item_but_retains_version(client):
    with Session(client.app.state.database_engine) as session:
        item = CandidateRepository(session).add_items(
            TEST_USER_ID, [SkillItemInput(name="Python", level="used")]
        )[0]
        item_id = item.id
        version_id = item.versions[0].id
        session.commit()

    assert [row["id"] for row in client.get("/api/v1/candidate-base/items").json()] == [str(item_id)]
    assert client.delete(f"/api/v1/candidate-base/items/{item_id}").status_code == 204
    assert client.get("/api/v1/candidate-base/items").json() == []
    with Session(client.app.state.database_engine) as session:
        stored = session.get(CandidateItemModel, item_id)
        assert stored is not None and stored.is_deleted and stored.deleted_at is not None
        assert session.get(CandidateItemVersionModel, version_id) is not None


def test_draft_review_and_delete_preserve_draft_blocks(client):
    created = client.post("/api/v1/resume-drafts")
    assert created.status_code == 201
    draft_id = created.json()["draft_id"]
    with_block = client.post(
        f"/api/v1/resume-drafts/{draft_id}/experience-blocks", json={"text": "QA"}
    ).json()
    block_id = with_block["blocks"][0]["id"]

    assert client.get("/api/v1/resume-drafts").status_code == 200
    reviewed = client.post(f"/api/v1/resume-drafts/{draft_id}/review")
    assert reviewed.status_code == 200
    assert reviewed.json()["state"] == "reviewed"
    assert any(row["draft_id"] == draft_id for row in client.get("/api/v1/resume-drafts").json())
    assert any(
        row["draft_id"] == draft_id
        for row in client.get("/api/v1/resume-drafts?state=reviewed").json()
    )
    assert all(
        row["draft_id"] != draft_id
        for row in client.get("/api/v1/resume-drafts?state=needs_user_review").json()
    )
    assert client.get(f"/api/v1/resume-drafts/{draft_id}").status_code == 200

    assert client.delete(f"/api/v1/resume-drafts/{draft_id}").status_code == 204
    assert client.get(f"/api/v1/resume-drafts/{draft_id}").status_code == 404
    assert client.patch(
        f"/api/v1/resume-drafts/{draft_id}/blocks/{block_id}", json={"text": "Changed"}
    ).status_code == 404
    with Session(client.app.state.database_engine) as session:
        stored = session.get(ResumeDraftModel, UUID(draft_id))
        assert stored is not None and stored.is_deleted and stored.deleted_at is not None
        assert session.get(DraftBlockModel, UUID(block_id)) is not None


def test_applying_block_after_review_does_not_reopen_draft(client):
    created = client.post("/api/v1/resume-drafts").json()
    block = client.post(
        f"/api/v1/resume-drafts/{created['draft_id']}/experience-blocks", json={"text": "QA"}
    ).json()["blocks"][0]
    assert client.post(f"/api/v1/resume-drafts/{created['draft_id']}/review").status_code == 200
    applied = client.post(
        f"/api/v1/resume-drafts/{created['draft_id']}/blocks/{block['id']}/apply",
        json={"items": [{"kind": "skill", "name": "Testing"}], "idempotency_key": "reviewed-block"},
    )
    assert applied.status_code == 200
    assert client.get(f"/api/v1/resume-drafts/{created['draft_id']}").json()["state"] == "reviewed"
